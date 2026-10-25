# pantallas/importar_excel.py
# Importar / exportar con Excel (parecido a Odoo), para productos y categorías:
#   - Se descarga una plantilla (vacía, con todo o solo lo seleccionado).
#   - Se edita en Excel: filas CON ID actualizan ese registro,
#     filas SIN ID crean registros nuevos.
#   - Al subirla se muestra una vista previa y solo se importa si no hay errores.
import os
import unicodedata
import tkinter as tk
import ttkbootstrap as ttk
from ttkbootstrap.constants import *
from tkinter import messagebox, filedialog
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import estilos
from conexion import supabase
from componentes.boton import Boton
from pantallas.categorias import (cargar_categorias, ruta_de, partir_ruta, buscar_ruta,
                                  buscar_hija, crear_ruta, con_descendientes,
                                  rutas_ordenadas, preparar_dialogo, SEPARADOR)

# campo en Supabase -> título de la columna en Excel (en este orden)
COLUMNAS = {
    "id": "ID",
    "nombre": "Nombre",
    "codigos": "Códigos de barras",
    "categoria_id": "Categoría",
    "unidad_base": "Unidad base",
    "precio_costo": "Precio costo",
    "precio_venta": "Precio venta",
    "stock": "Stock",
    "stock_minimo": "Stock mínimo",
}
NUMERICOS = ("precio_costo", "precio_venta", "stock", "stock_minimo")
ANCHOS = {"id": 8, "nombre": 40, "codigos": 30, "categoria_id": 30, "unidad_base": 14,
          "precio_costo": 14, "precio_venta": 14, "stock": 10, "stock_minimo": 14}

COLUMNAS_CAT = {"id": "ID", "nombre": "Nombre", "padre": "Categoría padre",
                "ruta": "Ruta completa (solo referencia)"}
ANCHOS_CAT = {"id": 8, "nombre": 30, "padre": 40, "ruta": 50}

FILAS_FORMATO_TEXTO = 500   # filas de la plantilla vacía con códigos como texto
EJEMPLO = "EJEMPLO"         # en la columna ID: el importador ignora esas filas

# filas de ejemplo de las plantillas vacías
EJEMPLOS_PRODUCTOS = [
    [EJEMPLO, "Gaseosa Inca Kola 500ml", "7751271001234",
     "Abarrotes / Bebidas / Gaseosas", "unidad", 1.8, 2.5, 24, 6],
    [EJEMPLO, "Arroz Costeño 5kg", "7750243001111, 7750243001128",
     "Abarrotes / Granos", "unidad", 16, 19.9, 10, 3],
]
EJEMPLOS_CATEGORIAS = [
    [EJEMPLO, "Abarrotes", "", "Abarrotes"],
    [EJEMPLO, "Bebidas", "Abarrotes", "Abarrotes / Bebidas"],
    [EJEMPLO, "Gaseosas", "Abarrotes / Bebidas", "Abarrotes / Bebidas / Gaseosas"],
]


def normalizar(titulo):
    """'Códigos de barras' -> 'codigosdebarras' (sin tildes, espacios ni mayúsculas)."""
    texto = unicodedata.normalize("NFKD", str(titulo or ""))
    return "".join(c for c in texto.lower() if c.isalnum())


# nombres de columna aceptados al leer el archivo
ALIAS = {normalizar(titulo): campo for campo, titulo in COLUMNAS.items()}
ALIAS.update({
    "codigo": "codigos", "codigos": "codigos", "codigodebarras": "codigos",
    "unidad": "unidad_base", "preciodecosto": "precio_costo",
    "preciodeventa": "precio_venta", "stockmin": "stock_minimo",
    "stockminimo": "stock_minimo", "categorias": "categoria_id",
    "familia": "categoria_id",
})
# la ruta completa no se lee: solo ayuda a ver el árbol en Excel
ALIAS_CAT = {"id": "id", "nombre": "nombre", "categoriapadre": "padre",
             "padre": "padre", "dentrode": "padre"}


def texto_celda(valor):
    """Valor de una celda como texto. 7501234567890.0 -> '7501234567890'"""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def numero_celda(valor):
    """Celda vacía -> 0. Acepta coma decimal ('2,5')."""
    if valor is None or texto_celda(valor) == "":
        return 0.0
    if isinstance(valor, (int, float)):
        return float(valor)
    return float(texto_celda(valor).replace(",", "."))


def separar_codigos(valor):
    """'111, 222; 333' -> ['111', '222', '333'] sin repetidos."""
    texto = texto_celda(valor).replace(";", ",").replace("\n", ",")
    codigos = []
    for c in texto.split(","):
        c = c.strip()
        if c and c not in codigos:
            codigos.append(c)
    return codigos


def codigos_de(producto):
    return [c["codigo"] for c in producto.get("codigos_barras") or []]


def valor_excel(valor):
    """100.0 -> 100 para que en Excel no salga '100.0'."""
    valor = float(valor or 0)
    return int(valor) if valor.is_integer() else valor


# ─── Exportar ────────────────────────────────

def guardar_libro(ruta, nombre_hoja, columnas, anchos, filas, instrucciones,
                  columnas_texto=(), ejemplos=(), calculadas=()):
    """Escribe un .xlsx con encabezado morado y una hoja de instrucciones.
    ejemplos: filas de muestra (en gris) que el importador ignora.
    calculadas: [(título, ancho, fórmula con {n} = n° de fila, formato)] que van
    al final; Excel las recalcula solo y el importador no las lee."""
    wb = Workbook()
    hoja = wb.active
    hoja.title = nombre_hoja
    hoja.append(list(columnas.values()) + [c[0] for c in calculadas])
    for fila in ejemplos:
        hoja.append(fila)
        for celda in hoja[hoja.max_row]:
            celda.font = Font(italic=True, color="8A8A9A")
    for fila in filas:
        hoja.append(fila)

    relleno = PatternFill("solid", fgColor=estilos.COLOR_SIDEBAR.lstrip("#"))
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno
        celda.alignment = Alignment(horizontal="center")
    for i, campo in enumerate(columnas, start=1):
        hoja.column_dimensions[hoja.cell(1, i).column_letter].width = anchos[campo]
    gris = PatternFill("solid", fgColor="E4E4EE")
    for i, (_, ancho, formula, formato) in enumerate(calculadas, start=len(columnas) + 1):
        hoja.cell(1, i).fill = PatternFill("solid", fgColor="8B7FD9")   # otro tono: no se edita
        hoja.column_dimensions[hoja.cell(1, i).column_letter].width = ancho
        for n in range(2, max(len(filas) + len(ejemplos), FILAS_FORMATO_TEXTO) + 2):
            celda = hoja.cell(n, i, formula.format(n=n))
            celda.number_format = formato
            celda.fill = gris
    # columnas como texto: si no, Excel vuelve los códigos 7.50123E+12
    for campo in columnas_texto:
        col = list(columnas).index(campo) + 1
        for n in range(2, max(len(filas) + len(ejemplos), FILAS_FORMATO_TEXTO) + 2):
            hoja.cell(n, col).number_format = "@"
    hoja.freeze_panes = "A2"

    ayuda = wb.create_sheet("Instrucciones")
    ayuda.column_dimensions["A"].width = 100
    if ejemplos:
        instrucciones = (f"• Las filas en gris con {EJEMPLO} en la columna ID son de muestra: "
                         "no se importan. Puedes borrarlas o dejarlas.",) + tuple(instrucciones)
    for linea in ("CÓMO USAR ESTA PLANTILLA", "") + tuple(instrucciones):
        ayuda.append([linea])
    ayuda["A1"].font = Font(bold=True, size=13)
    wb.save(ruta)


def exportar_plantilla(ruta, productos, cats=None):
    """Plantilla de productos. Con productos = lista de Supabase, sin = vacía
    (la vacía lleva filas de ejemplo).
    cats: categorías (id -> fila) para escribir su nombre completo."""
    cats = cats or {}
    filas = [[p["id"], p.get("nombre", ""), ", ".join(codigos_de(p)),
              ruta_de(p.get("categoria_id"), cats), p.get("unidad_base") or "unidad",
              *(valor_excel(p.get(campo)) for campo in NUMERICOS)]
             for p in productos]
    guardar_libro(ruta, "Productos", COLUMNAS, ANCHOS, filas, (
        "• Filas CON ID: se actualiza ese producto con los datos de la fila.",
        "• Filas SIN ID (vacío): se crea un producto nuevo.",
        "• No cambies los números de ID ni los títulos de la primera fila.",
        "• Nombre es obligatorio.",
        "• Códigos de barras: separados por coma. Ej: 7501234567890, 7509876543210",
        "  Un código no puede repetirse ni pertenecer a otro producto.",
        "• Categoría: nombre completo con sus subcategorías separadas por /.",
        "  Ej: Abarrotes / Bebidas / Gaseosas",
        "  Si no existe, se crea sola al importar (con todas sus subcategorías).",
        "  Vacío = sin categoría.",
        "• Precios y stock deben ser números (vacío = 0), en soles.",
        "• El precio de costo no puede ser mayor que el precio de venta.",
        "• Las columnas grises del final se calculan solas (no se importan):",
        "    Ganancia S/ = lo que ganas por cada unidad (venta - costo).",
        "    Margen %    = de cada S/ 1.00 que cobras, cuánto es ganancia.",
        "                  Ej: 28% significa que de cada S/ 1.00, S/ 0.28 es para ti.",
        "• Puedes borrar columnas que no quieras modificar (excepto ID y Nombre).",
        "• Quitar una fila del archivo NO elimina el producto.",
    ), columnas_texto=("codigos",), ejemplos=[] if productos else EJEMPLOS_PRODUCTOS,
       calculadas=formulas_margen())


def formulas_margen():
    """Columnas Ganancia y Margen % calculadas con fórmulas de Excel."""
    letra = {campo: get_column_letter(i) for i, campo in enumerate(COLUMNAS, start=1)}
    c, v = letra["precio_costo"], letra["precio_venta"]
    return [
        ("Ganancia S/", 13, f'=IF({v}{{n}}="","",{v}{{n}}-{c}{{n}})', '"S/ "#,##0.00'),
        ("Margen % (de cada S/1 cobrado)", 30, f'=IF(N({v}{{n}})>0,({v}{{n}}-{c}{{n}})/{v}{{n}},"")', "0.0%"),
    ]


def exportar_categorias(ruta, cats):
    """Plantilla de categorías (en orden de árbol). cats vacío = plantilla vacía."""
    filas = [[cid, cats[cid]["nombre"], ruta_de(cats[cid]["padre_id"], cats), texto]
             for texto, cid in rutas_ordenadas(cats)]
    guardar_libro(ruta, "Categorías", COLUMNAS_CAT, ANCHOS_CAT, filas, (
        "• Filas CON ID: se actualiza esa categoría (nombre o categoría padre).",
        "• Filas SIN ID (vacío): se crea una categoría nueva.",
        "• Nombre: solo el nombre de esta categoría, sin /. Ej: Gaseosas",
        "• Categoría padre: dentro de qué categoría va, con la ruta completa.",
        "  Ej: Abarrotes / Bebidas      (vacío = categoría principal)",
        "  Si la categoría padre no existe, se crea sola al importar.",
        "• Puedes anidar todos los niveles que quieras. Ejemplo:",
        "      Nombre: Abarrotes   Categoría padre: (vacío)",
        "      Nombre: Bebidas     Categoría padre: Abarrotes",
        "      Nombre: Gaseosas    Categoría padre: Abarrotes / Bebidas",
        "• La columna 'Ruta completa' es solo para ver el árbol; no se lee al importar.",
        "• Si una fila sin ID ya existe, se deja como está (no se duplica).",
        "• Quitar una fila del archivo NO elimina la categoría.",
    ), ejemplos=[] if cats else EJEMPLOS_CATEGORIAS)


def guardar_excel(dueno, nombre_archivo, escribir):
    """Pregunta dónde guardar, escribe el archivo y ofrece abrirlo."""
    ruta = filedialog.asksaveasfilename(
        parent=dueno, title="Guardar Excel", initialfile=nombre_archivo,
        defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")])
    if not ruta:
        return
    try:
        escribir(ruta)
    except PermissionError:
        messagebox.showerror("Error", "No se pudo guardar el archivo.\n"
                             "Si está abierto en Excel, ciérralo e intenta de nuevo.",
                             parent=dueno)
        return
    except Exception as e:
        messagebox.showerror("Error", "No se pudo exportar:\n" + str(e), parent=dueno)
        return
    if messagebox.askyesno("Listo", "Archivo guardado.\n\n¿Abrirlo ahora en Excel?",
                           parent=dueno):
        try:
            os.startfile(ruta)
        except Exception as e:
            messagebox.showerror("Error", "No se pudo abrir:\n" + str(e), parent=dueno)


def exportar_productos(dueno, productos):
    """Exporta esa lista de productos (p. ej. los seleccionados en la tabla)."""
    guardar_excel(dueno, "productos.xlsx",
                  lambda ruta: exportar_plantilla(ruta, productos, cargar_categorias()))


# ─── Leer ────────────────────────────────────

def leer_excel(ruta, alias):
    """Devuelve (columnas presentes, [(n° de fila, {campo: valor})]) sin filas vacías."""
    wb = load_workbook(ruta, read_only=True, data_only=True)
    filas_excel = list(wb.worksheets[0].iter_rows(values_only=True))
    wb.close()
    if not filas_excel:
        raise ValueError("El archivo está vacío.")
    mapa = {}   # índice de columna -> campo
    for i, titulo in enumerate(filas_excel[0]):
        campo = alias.get(normalizar(titulo))
        if campo and campo not in mapa.values():
            mapa[i] = campo
    filas = []
    for n, valores in enumerate(filas_excel[1:], start=2):
        celdas = {campo: valores[i] if i < len(valores) else None
                  for i, campo in mapa.items()}
        if texto_celda(celdas.get("id")).upper() == EJEMPLO:
            continue   # fila de muestra de la plantilla
        if any(texto_celda(v) != "" for v in celdas.values()):
            filas.append((n, celdas))
    return set(mapa.values()), filas


def leer_id(id_txt, existentes, ids_vistos, n, errores, que="un producto"):
    """Valida la celda ID. Devuelve el id (int) o None."""
    if not id_txt:
        return None
    if not id_txt.isdigit():
        errores.append(f"ID '{id_txt}' no es válido")
    elif int(id_txt) not in existentes:
        errores.append(f"No existe {que} con ID {id_txt}")
    elif int(id_txt) in ids_vistos:
        errores.append(f"ID {id_txt} repetido (fila {ids_vistos[int(id_txt)]})")
    else:
        ids_vistos[int(id_txt)] = n
        return int(id_txt)
    return None


def poner_accion(fila, hay_cambios):
    if fila["errores"]:
        fila["accion"] = "error"
        fila["detalle"] = " · ".join(fila["errores"])
    elif fila["id"] is None:
        fila["accion"] = "nuevo"
    else:
        fila["accion"] = "actualizar" if hay_cambios else "igual"


# ─── Productos: validar y guardar ────────────

def analizar_archivo(ruta):
    """Lee el Excel de productos y compara con Supabase. Devuelve una lista de filas
    con 'accion' = nuevo | actualizar | igual | error. No guarda nada."""
    presentes, filas_excel = leer_excel(ruta, ALIAS)
    if not presentes & {"id", "nombre"}:
        raise ValueError("No se encontraron las columnas ID ni Nombre en la primera fila.\n"
                         "Usa la plantilla de productos descargada desde este programa.")

    resp = supabase.table("productos").select("*, codigos_barras(codigo)").execute()
    existentes = {p["id"]: p for p in resp.data}
    dueno_db = {c: p for p in resp.data for c in codigos_de(p)}
    cats = cargar_categorias()

    resultado = []
    ids_vistos = {}   # id -> fila de Excel donde apareció
    for n, celdas in filas_excel:
        errores = []
        fila = {"fila": n, "datos": {}, "codigos": None, "codigos_antes": [],
                "errores": errores, "categoria_nueva": None, "detalle": "",
                "nombre": texto_celda(celdas.get("nombre"))}
        resultado.append(fila)

        id_txt = texto_celda(celdas.get("id"))
        fila["id"] = leer_id(id_txt, existentes, ids_vistos, n, errores)
        actual = existentes.get(fila["id"], {})
        fila["codigos_antes"] = codigos_de(actual)

        datos = fila["datos"]
        if "nombre" in presentes:
            if not fila["nombre"]:
                errores.append("El nombre es obligatorio")
            datos["nombre"] = fila["nombre"]
        elif not id_txt:
            errores.append("Falta la columna Nombre para crear productos")
        fila["nombre"] = fila["nombre"] or actual.get("nombre", "")
        if "unidad_base" in presentes:
            datos["unidad_base"] = texto_celda(celdas["unidad_base"]) or "unidad"
        for campo in NUMERICOS:
            if campo in presentes:
                try:
                    datos[campo] = numero_celda(celdas[campo])
                except ValueError:
                    errores.append(f"{COLUMNAS[campo]}: '{texto_celda(celdas[campo])}' "
                                   "no es un número")
        if "codigos" in presentes:
            fila["codigos"] = separar_codigos(celdas["codigos"])
        if "categoria_id" in presentes:
            partes = partir_ruta(texto_celda(celdas["categoria_id"]))
            datos["categoria_id"] = buscar_ruta(partes, cats) if partes else None
            if partes and datos["categoria_id"] is None:
                fila["categoria_nueva"] = SEPARADOR.join(partes)   # se crea al importar

        # costo <= venta (si falta una columna se usa el valor que ya tiene)
        costo = datos.get("precio_costo", float(actual.get("precio_costo") or 0))
        venta = datos.get("precio_venta", float(actual.get("precio_venta") or 0))
        if costo > venta:
            errores.append(f"El precio de costo ({valor_excel(costo)}) no puede ser "
                           f"mayor que el de venta ({valor_excel(venta)})")

        if not id_txt:
            # producto nuevo: completar lo que falte
            datos.setdefault("unidad_base", "unidad")
            datos.setdefault("categoria_id", None)
            for campo in NUMERICOS:
                datos.setdefault(campo, 0.0)
            datos["activo"] = True
        elif actual:
            # actualización: solo lo que cambió
            fila["datos"] = {k: v for k, v in datos.items()
                             if not igual(actual.get(k), v)
                             or (k == "categoria_id" and fila["categoria_nueva"])}

    validar_codigos(resultado, dueno_db)

    for fila in resultado:
        cambios = [COLUMNAS[k] for k in fila["datos"]] if fila["id"] else []
        if cambian_codigos(fila):
            cambios.append(COLUMNAS["codigos"])
        poner_accion(fila, bool(cambios))
        if fila["accion"] == "actualizar":
            fila["detalle"] = "Cambia: " + ", ".join(cambios)
        if fila["categoria_nueva"] and fila["accion"] != "error":
            fila["detalle"] += (" · " if fila["detalle"] else "") + \
                f"Crea categoría: {fila['categoria_nueva']}"
    return resultado


def igual(antes, despues):
    if isinstance(despues, float):
        return abs(float(antes or 0) - despues) < 1e-9
    if isinstance(despues, str):
        return (antes or "") == despues
    return antes == despues   # categoria_id: número o None


def cambian_codigos(fila):
    return fila["codigos"] is not None and set(fila["codigos"]) != set(fila["codigos_antes"])


def validar_codigos(resultado, dueno_db):
    """Un código no puede estar dos veces en el archivo ni pertenecer a otro
    producto, salvo que el archivo se lo quite a ese producto."""
    liberados = set()
    for fila in resultado:
        if fila["id"] is not None and fila["codigos"] is not None:
            liberados |= set(fila["codigos_antes"]) - set(fila["codigos"])
    vistos = {}
    for fila in resultado:
        for c in fila["codigos"] or []:
            if c in vistos:
                fila["errores"].append(f"Código {c} repetido (fila {vistos[c]})")
                continue
            vistos[c] = fila["fila"]
            dueno = dueno_db.get(c)
            if dueno and dueno["id"] != fila["id"] and c not in liberados:
                fila["errores"].append(f"El código {c} ya es del producto "
                                       f"#{dueno['id']} - {dueno.get('nombre', '')}")


def importar_filas(resultado):
    """Guarda en Supabase los productos nuevos y los que cambian."""
    actualizar = [f for f in resultado if f["accion"] == "actualizar"]
    nuevos = [f for f in resultado if f["accion"] == "nuevo"]

    # 1. quitar primero los códigos que salen, por si pasan a otro producto
    for f in actualizar:
        if f["codigos"] is not None:
            quitados = [c for c in f["codigos_antes"] if c not in f["codigos"]]
            if quitados:
                supabase.table("codigos_barras").delete() \
                    .eq("producto_id", f["id"]).in_("codigo", quitados).execute()
    # 2. crear las categorías que no existían (con sus subcategorías)
    cats = cargar_categorias()
    for f in actualizar + nuevos:
        if f["categoria_nueva"]:
            f["datos"]["categoria_id"] = crear_ruta(partir_ruta(f["categoria_nueva"]), cats)
    # 3. actualizar productos existentes
    for f in actualizar:
        if f["datos"]:
            supabase.table("productos").update(f["datos"]).eq("id", f["id"]).execute()
    # 4. crear productos nuevos (todos de una vez; vuelven en el mismo orden)
    if nuevos:
        resp = supabase.table("productos").insert([f["datos"] for f in nuevos]).execute()
        for f, p in zip(nuevos, resp.data):
            f["id"] = p["id"]
    # 5. códigos nuevos de todos
    codigos = [{"producto_id": f["id"], "codigo": c}
               for f in actualizar + nuevos
               for c in f["codigos"] or [] if c not in f["codigos_antes"]]
    if codigos:
        supabase.table("codigos_barras").insert(codigos).execute()
    return len(nuevos), len(actualizar)


# ─── Categorías: validar y guardar ───────────

def mas_profunda_existente(partes, cats):
    """De 'A / B / C', la última parte que ya existe (su id) o None."""
    padre_id = None
    for nombre in partes:
        cat = buscar_hija(padre_id, nombre, cats)
        if cat is None:
            break
        padre_id = cat["id"]
    return padre_id


def analizar_categorias(ruta):
    """Lee el Excel de categorías y compara con Supabase. No guarda nada."""
    presentes, filas_excel = leer_excel(ruta, ALIAS_CAT)
    if "nombre" not in presentes:
        raise ValueError("No se encontró la columna Nombre en la primera fila.\n"
                         "Usa la plantilla de categorías descargada desde este programa.")
    cats = cargar_categorias()

    resultado = []
    ids_vistos = {}
    rutas_vistas = {}   # ruta final en minúsculas -> fila de Excel
    for n, celdas in filas_excel:
        errores = []
        nombre = " ".join(texto_celda(celdas.get("nombre")).split())
        fila = {"fila": n, "nombre": nombre, "padre_id": None, "padre_nuevo": None,
                "errores": errores, "detalle": ""}
        resultado.append(fila)
        fila["id"] = leer_id(texto_celda(celdas.get("id")), cats, ids_vistos, n,
                             errores, "una categoría")
        actual = cats.get(fila["id"], {})

        if not nombre:
            errores.append("El nombre es obligatorio")
        elif "/" in nombre:
            errores.append("El nombre no puede llevar / (la jerarquía va en Categoría padre)")

        if "padre" in presentes:
            partes = partir_ruta(texto_celda(celdas.get("padre")))
        else:   # sin la columna se queda donde está
            partes = partir_ruta(ruta_de(actual.get("padre_id"), cats))
        if partes:
            fila["padre_id"] = buscar_ruta(partes, cats)
            if fila["padre_id"] is None:
                fila["padre_nuevo"] = SEPARADOR.join(partes)   # se crea al importar
        if fila["id"] is not None and \
                mas_profunda_existente(partes, cats) in con_descendientes(fila["id"], cats):
            errores.append("No puede ir dentro de sí misma ni de sus subcategorías")

        ruta_padre = fila["padre_nuevo"] or ruta_de(fila["padre_id"], cats)
        fila["ruta"] = ruta_padre + SEPARADOR + nombre if ruta_padre else nombre
        clave = fila["ruta"].lower()
        if nombre and clave in rutas_vistas:
            errores.append(f"Repetida (fila {rutas_vistas[clave]})")
        rutas_vistas.setdefault(clave, n)

        existente = None if fila["padre_nuevo"] or not nombre else \
            buscar_hija(fila["padre_id"], nombre, cats)
        cambios = []
        if fila["id"] is None:
            if existente and not errores:
                # ya existe: no se duplica (así se puede importar el mismo archivo dos veces)
                fila["accion"], fila["detalle"] = "igual", "Ya existe"
                fila["id"] = existente["id"]
                continue
        elif actual:
            if existente and existente["id"] != fila["id"]:
                errores.append(f"Ya existe otra categoría \"{fila['ruta']}\"")
            if nombre != actual["nombre"]:
                cambios.append("Nombre")
            if fila["padre_nuevo"] or fila["padre_id"] != actual["padre_id"]:
                cambios.append("Categoría padre")
        poner_accion(fila, bool(cambios))
        if fila["accion"] == "actualizar":
            fila["detalle"] = "Cambia: " + ", ".join(cambios)
        if fila["padre_nuevo"] and fila["accion"] != "error":
            fila["detalle"] += (" · " if fila["detalle"] else "") + \
                f"Crea: {fila['padre_nuevo']}"

    for fila in resultado:
        fila["nombre_mostrar"] = fila["ruta"]   # en la vista previa se ve el árbol
    return resultado


def importar_categorias(resultado):
    """Guarda en Supabase las categorías nuevas y las que cambian (en orden)."""
    cats = cargar_categorias()
    antes = len(cats)
    actualizadas = 0
    for f in resultado:
        if f["accion"] not in ("nuevo", "actualizar"):
            continue
        padre_id = crear_ruta(partir_ruta(f["padre_nuevo"]), cats) \
            if f["padre_nuevo"] else f["padre_id"]
        datos = {"nombre": f["nombre"], "padre_id": padre_id}
        if f["accion"] == "nuevo":
            if buscar_hija(padre_id, f["nombre"], cats):
                continue   # ya la creó otra fila como categoría padre
            cat = supabase.table("categorias").insert(datos).execute().data[0]
            cats[cat["id"]] = cat
        else:
            supabase.table("categorias").update(datos).eq("id", f["id"]).execute()
            cats[f["id"]].update(datos)
            actualizadas += 1
    return len(cats) - antes, actualizadas


# ─── Ventanas ────────────────────────────────

def abrir_importador(padre, al_importar, seleccionados=()):
    """Importar productos. seleccionados: productos marcados en la tabla."""
    exportaciones = []
    if seleccionados:
        exportaciones.append((f"Exportar seleccionados ({len(seleccionados)})",
                              lambda dueno: exportar_productos(dueno, list(seleccionados))))

    def exportar_todos(dueno):
        productos = (supabase.table("productos").select("*, codigos_barras(codigo)")
                     .order("id").execute()).data
        exportar_productos(dueno, productos)

    exportaciones += [
        ("Exportar todos los productos", exportar_todos),
        ("Plantilla vacía", lambda dueno: guardar_excel(
            dueno, "plantilla_productos.xlsx", lambda r: exportar_plantilla(r, []))),
    ]
    abrir_dialogo_importar(
        padre, "Importar productos desde Excel",
        "Filas con ID actualizan ese producto · filas sin ID crean productos nuevos · "
        "categoría con ruta completa: Abarrotes / Bebidas",
        exportaciones, analizar_archivo, importar_filas, al_importar,
        ("Productos nuevos", "Productos actualizados"),
        "Algún código de barras ya existe en otro producto.")


def abrir_importador_categorias(padre, al_importar):
    """Importar categorías y subcategorías."""
    abrir_dialogo_importar(
        padre, "Importar categorías desde Excel",
        "Filas con ID actualizan esa categoría · filas sin ID crean categorías nuevas · "
        "Categoría padre con ruta completa: Abarrotes / Bebidas",
        [("Exportar todas las categorías", lambda dueno: guardar_excel(
            dueno, "categorias.xlsx", lambda r: exportar_categorias(r, cargar_categorias()))),
         ("Plantilla vacía", lambda dueno: guardar_excel(
             dueno, "plantilla_categorias.xlsx", lambda r: exportar_categorias(r, {})))],
        analizar_categorias, importar_categorias, al_importar,
        ("Categorías nuevas", "Categorías actualizadas"),
        "Alguna categoría ya existe.")


def abrir_dialogo_importar(padre, titulo, ayuda, exportaciones, analizar, importar,
                           al_importar, etiquetas, aviso_duplicado):
    """Ventana común: paso 1 exportar plantilla, paso 2 elegir archivo y vista previa.
    exportaciones: [(texto del botón, función(dueño))]; el primero es el principal."""
    ventana = tk.Toplevel(padre)
    preparar_dialogo(ventana, padre, titulo, 900, 600)

    # el pie se empaqueta antes que el cuerpo para que la tabla nunca lo tape
    pie = tk.Frame(ventana, bg=estilos.COLOR_FONDO)
    pie.pack(side="bottom", fill="x", padx=25, pady=15)

    cuerpo = tk.Frame(ventana, bg=estilos.COLOR_FONDO)
    cuerpo.pack(fill="both", expand=True, padx=25, pady=(15, 0))
    cuerpo.grid_columnconfigure(0, weight=1)
    cuerpo.grid_rowconfigure(3, weight=1)   # la vista previa crece

    # ── Paso 1: plantilla ──
    paso1 = tk.Frame(cuerpo, bg="white", highlightthickness=1,
                     highlightbackground="#d0d0dc")
    paso1.grid(row=0, column=0, sticky="ew")
    tk.Label(paso1, text="1. Descarga la plantilla y llénala en Excel",
             font=estilos.FUENTE_BOTON, bg="white", fg=estilos.COLOR_TEXTO
             ).pack(anchor="w", padx=15, pady=(12, 2))
    tk.Label(paso1, text=ayuda, font=estilos.FUENTE_CARD_TXT, bg="white", fg="#8a8a9a"
             ).pack(anchor="w", padx=15)
    botones1 = tk.Frame(paso1, bg="white")
    botones1.pack(anchor="w", padx=15, pady=12)
    for i, (texto, accion) in enumerate(exportaciones):
        Boton(botones1, texto, lambda a=accion: exportar(a),
              estilo="primario" if i == 0 else "secundario"
              ).pack(side="left", padx=(0 if i == 0 else 8, 0))

    # ── Paso 2: subir ──
    paso2 = tk.Frame(cuerpo, bg=estilos.COLOR_FONDO)
    paso2.grid(row=1, column=0, sticky="ew", pady=(15, 8))
    tk.Label(paso2, text="2. Sube el archivo", font=estilos.FUENTE_BOTON,
             bg=estilos.COLOR_FONDO, fg=estilos.COLOR_TEXTO).pack(side="left")
    Boton(paso2, "Elegir archivo...", lambda: elegir_archivo(), estilo="secundario",
          pady=4).pack(side="left", padx=12)
    lbl_archivo = tk.Label(paso2, text="Ningún archivo seleccionado",
                           font=estilos.FUENTE_CARD_TXT, bg=estilos.COLOR_FONDO, fg="#8a8a9a")
    lbl_archivo.pack(side="left")

    lbl_resumen = tk.Label(cuerpo, text="", font=estilos.FUENTE_CARD_TXT,
                           bg=estilos.COLOR_FONDO, fg=estilos.COLOR_TEXTO, anchor="w")
    lbl_resumen.grid(row=2, column=0, sticky="ew", pady=(0, 5))

    # ── Vista previa ──
    previa_frame = tk.Frame(cuerpo, bg=estilos.COLOR_FONDO)
    previa_frame.grid(row=3, column=0, sticky="nsew")
    previa_frame.grid_columnconfigure(0, weight=1)
    previa_frame.grid_rowconfigure(0, weight=1)
    columnas = ("fila", "accion", "id", "nombre", "detalle")
    previa = ttk.Treeview(previa_frame, columns=columnas, show="headings",
                          selectmode="browse", height=6)
    for col, texto, ancho_col, ancla in (("fila", "Fila", 50, "center"),
                                        ("accion", "Acción", 100, "center"),
                                        ("id", "ID", 60, "center"),
                                        ("nombre", "Nombre", 240, "w"),
                                        ("detalle", "Detalle", 360, "w")):
        previa.heading(col, text=texto)
        previa.column(col, width=ancho_col, anchor=ancla, stretch=(col == "detalle"))
    previa.tag_configure("error", foreground=estilos.COLOR_PELIGRO)
    previa.tag_configure("nuevo", foreground=estilos.COLOR_EXITO)
    previa.tag_configure("actualizar", foreground=estilos.COLOR_SIDEBAR)
    previa.tag_configure("igual", foreground="#8a8a9a")
    scroll = ttk.Scrollbar(previa_frame, orient="vertical", command=previa.yview)
    previa.configure(yscrollcommand=scroll.set)
    previa.grid(row=0, column=0, sticky="nsew")
    scroll.grid(row=0, column=1, sticky="ns")

    # ── Pie ──
    btn_importar = Boton(pie, "Importar", lambda: guardar())
    btn_importar.pack(side="right")
    Boton(pie, "Cancelar", ventana.destroy, estilo="secundario").pack(side="right", padx=8)
    btn_importar.activar(False)

    estado = {"filas": []}
    NOMBRE_ACCION = {"nuevo": "Nuevo", "actualizar": "Actualizar",
                     "igual": "Sin cambios", "error": "Error"}

    def exportar(accion):
        try:
            accion(ventana)
        except Exception as e:
            messagebox.showerror("Error", "No se pudo exportar:\n" + str(e), parent=ventana)

    def elegir_archivo():
        ruta = filedialog.askopenfilename(
            parent=ventana, title="Elegir archivo de Excel",
            filetypes=[("Excel", "*.xlsx *.xlsm"), ("Todos los archivos", "*.*")])
        if not ruta:
            return
        lbl_archivo.configure(text=os.path.basename(ruta), fg=estilos.COLOR_TEXTO)
        previa.delete(*previa.get_children())
        estado["filas"] = []
        try:
            estado["filas"] = analizar(ruta)
        except PermissionError:
            mostrar_error("No se pudo leer el archivo. Si está abierto en Excel, "
                          "ciérralo e intenta de nuevo.")
            return
        except Exception as e:
            mostrar_error("No se pudo leer el archivo:\n" + str(e))
            return
        mostrar_previa()

    def mostrar_error(texto):
        lbl_resumen.configure(text="", fg=estilos.COLOR_TEXTO)
        btn_importar.activar(False)
        messagebox.showerror("Error", texto, parent=ventana)

    def mostrar_previa():
        cuenta = {accion: 0 for accion in NOMBRE_ACCION}
        for f in estado["filas"]:
            cuenta[f["accion"]] += 1
            previa.insert("", END, tags=(f["accion"],), values=(
                f["fila"], NOMBRE_ACCION[f["accion"]], f["id"] or "-",
                f.get("nombre_mostrar", f["nombre"]), f["detalle"]))

        if not estado["filas"]:
            lbl_resumen.configure(text="El archivo no tiene filas con datos.",
                                  fg=estilos.COLOR_PELIGRO)
        elif cuenta["error"]:
            lbl_resumen.configure(
                text=f"{cuenta['error']} fila(s) con errores. Corrígelas en Excel y "
                     "vuelve a elegir el archivo (no se importa nada mientras haya errores).",
                fg=estilos.COLOR_PELIGRO)
        else:
            lbl_resumen.configure(
                text=f"{cuenta['nuevo']} nuevos · {cuenta['actualizar']} a actualizar · "
                     f"{cuenta['igual']} sin cambios", fg=estilos.COLOR_TEXTO)
        btn_importar.activar(not cuenta["error"]
                             and cuenta["nuevo"] + cuenta["actualizar"] > 0)

    def guardar():
        try:
            nuevos, actualizados = importar(estado["filas"])
        except Exception as e:
            texto = str(e)
            if "duplicate" in texto.lower() or "23505" in texto:
                texto = aviso_duplicado + "\n\n" + texto
            messagebox.showerror("Error", "La importación se detuvo (puede haber "
                                 "quedado a medias):\n" + texto, parent=ventana)
            al_importar()
            return
        messagebox.showinfo("Listo", f"Importación terminada.\n\n"
                                     f"{etiquetas[0]}: {nuevos}\n"
                                     f"{etiquetas[1]}: {actualizados}",
                            parent=ventana)
        ventana.destroy()
        al_importar()   # recargar la tabla
